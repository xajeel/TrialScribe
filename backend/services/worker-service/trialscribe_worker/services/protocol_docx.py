"""Assemble a tenant-scoped ICH M11 Word document from stored chapters."""

from collections.abc import Sequence
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from docx import Document

from trialscribe_worker.models.evidence_chunk import EvidenceChunk
from trialscribe_worker.models.m11_section_record import M11SectionListRecord
from trialscribe_worker.retrieval.citations import (
    apply_citations,
    apply_numbered_citations,
    parse_cite_ids,
)
from trialscribe_worker.utils.constant import (
    EXPORT_SCOPE_DONE_ONLY,
    M11_SECTION_DONE_STATUS,
    MAX_EXPORT_FILENAME_LENGTH,
)

SCOPE_DONE_ONLY_LABEL = "Done sections only"
SCOPE_INCLUDE_DRAFTS_LABEL = "Includes unfinished sections (drafts marked)"
REFERENCES_HEADING = "References"
DRAFT_HEADING_SUFFIX = " (Draft)"


def export_filename(protocol_title: str, exported_at: datetime) -> str:
    """Return a download name derived from the protocol title and UTC date."""

    date = exported_at.astimezone(UTC).strftime("%Y-%m-%d")
    slug_chars: list[str] = []
    for character in protocol_title.strip():
        if character.isspace():
            slug_chars.append("_")
        elif character.isascii() and (character.isalnum() or character in "._-"):
            slug_chars.append(character)
    slug = "".join(slug_chars).strip("._-") or "protocol"
    suffix = f"_protocol_{date}.docx"
    limit = MAX_EXPORT_FILENAME_LENGTH - len(suffix)
    if limit < 1:
        return "protocol.docx"
    return f"{slug[:limit]}{suffix}"


def build_protocol_docx(
    *,
    protocol_title: str,
    exported_at: datetime,
    scope: str,
    sections: Sequence[M11SectionListRecord],
    chunks: Sequence[EvidenceChunk],
) -> bytes:
    """Return Word bytes for the included chapters. Never writes passage text."""

    included = _included_sections(scope, sections)
    allowed = {chunk.id for chunk in chunks}
    ordered = _citation_order(included, allowed)
    by_id = {chunk.id: chunk for chunk in chunks}

    document = Document()
    document.core_properties.title = protocol_title
    document.add_heading(protocol_title, level=0)
    document.add_paragraph(
        SCOPE_DONE_ONLY_LABEL
        if scope == EXPORT_SCOPE_DONE_ONLY
        else SCOPE_INCLUDE_DRAFTS_LABEL
    )
    document.add_paragraph(exported_at.astimezone(UTC).strftime("%Y-%m-%d"))

    for section in included:
        heading = f"{section.section_number}. {section.title}"
        if section.status != M11_SECTION_DONE_STATUS:
            heading += DRAFT_HEADING_SUFFIX
        document.add_heading(heading, level=1)
        body = apply_numbered_citations(section.content, ordered)
        for line in body.splitlines():
            if line.strip() != "":
                document.add_paragraph(line)

    if ordered:
        document.add_heading(REFERENCES_HEADING, level=1)
        for number, chunk_id in enumerate(ordered, start=1):
            chunk = by_id.get(chunk_id)
            if chunk is None:
                continue
            line = f"[{number}] {chunk.source_kind}: {chunk.source_identity}"
            if chunk.page_number is not None:
                line += f", p. {chunk.page_number}"
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _included_sections(
    scope: str,
    sections: Sequence[M11SectionListRecord],
) -> list[M11SectionListRecord]:
    if scope == EXPORT_SCOPE_DONE_ONLY:
        return [row for row in sections if row.status == M11_SECTION_DONE_STATUS]
    return list(sections)


def _citation_order(
    sections: Sequence[M11SectionListRecord],
    allowed: set[UUID],
) -> list[UUID]:
    ordered: list[UUID] = []
    seen: set[UUID] = set()
    for section in sections:
        cleaned = apply_citations(section.content, allowed)
        for chunk_id in parse_cite_ids(cleaned):
            if chunk_id in seen:
                continue
            seen.add(chunk_id)
            ordered.append(chunk_id)
    return ordered
